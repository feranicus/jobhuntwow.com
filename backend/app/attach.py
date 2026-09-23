"""attach.py - files that belong to ONE application: a transcript, a deck, a take-home task.

He asked for it in one line: *"maybe a good idea also to add a file like txt from transcript or ppt
also would be good"*. An interview happens, he has the transcript or the slides they sent, and the
only sane place for it is the card for that job.

THREE RULES, and every one of them is a defect class this repo has already paid for.

1. **THE FILE IS KEPT, AND SO IS ITS TEXT.** A `.pptx` on disk is bytes nobody can search. Every
   upload is read ONCE, at the moment it arrives, into plain text, and that text is stored beside
   it - so the transcript of an interview is searchable, and can be fed to the tailor later without
   re-parsing anything. The extraction is DETERMINISTIC (pypdf / python-docx / python-pptx / the
   bytes themselves). No model reads an upload.

2. **A FILE WITH NO TEXT SAYS SO.** A scanned PDF and a photo of a whiteboard extract to nothing.
   The file is still kept - it is his - but `text_note` records WHY there is no text, because
   "no text" and "we failed to read it" are different facts and they must not render the same.
   Absence of evidence is never a finding.

3. **THE NAME IS OURS, NOT THE BROWSER'S.** `filename` is attacker-controlled: it may contain
   `..`, a drive letter, a NUL, or a path separator. `safe_name()` reduces it to one path segment
   with an allowlisted extension, and a collision is NUMBERED rather than overwritten - an upload
   must never silently destroy a file he already has.

Storage is the SAME tenancy the documents already use (`DATA_DIR/users/<hash>/<job_id>/`), under
`attachments/`, so an attachment cannot outlive or escape the application it belongs to. The
INDEX is `tracker.attachments`; the bytes are on disk. One home for each.
"""
from __future__ import annotations

import io
import json
import os
import re
import time
import unicodedata

# What a person actually attaches to a job application. Deliberately SHORT: every extension here
# is inert data. No .html (script), no .svg (script), no archives, no executables - an allowlist
# is the only list that stays safe when a new file type appears in the world.
TEXT_EXT = {".txt", ".md", ".log", ".csv", ".vtt", ".srt", ".json"}
RICH_EXT = {".pdf", ".docx", ".pptx", ".xlsx"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp"}
ALLOWED_EXT = TEXT_EXT | RICH_EXT | IMAGE_EXT

MAX_BYTES = int(os.getenv("JHW_ATTACH_MAX_MB", "25")) * 1024 * 1024
MAX_TEXT = int(os.getenv("JHW_ATTACH_TEXT_CHARS", "200000"))
MAX_PER_JOB = int(os.getenv("JHW_ATTACH_MAX_PER_JOB", "40"))

_SAFE = re.compile(r"[^A-Za-z0-9._-]+")
_TIMECODE = re.compile(
    r"^\s*(\d+\s*$|\d{1,2}:\d{2}(:\d{2})?([.,]\d{1,3})?\s*-->\s*"
    r"\d{1,2}:\d{2}(:\d{2})?([.,]\d{1,3})?.*$|WEBVTT.*$|NOTE\s.*$)")


def _log(**k):
    k.setdefault("ts", time.time())
    k.setdefault("service", "jhw-web")
    try:
        print(json.dumps(k, ensure_ascii=False), flush=True)
    except Exception:
        pass


def ext_of(name: str) -> str:
    return os.path.splitext((name or "").strip().lower())[1]


def safe_name(name: str) -> str:
    """ONE path segment, allowlisted extension, never empty, never a traversal.

    `os.path.basename` alone is not enough: it leaves a Windows `C:\\x\\y` intact on POSIX, so the
    separators are normalised FIRST. A NUL byte truncates a path inside the C library, so it goes
    too. The result is what we store - the browser's string is never used as a path."""
    raw = (name or "").replace("\x00", "")
    raw = raw.replace("\\", "/").split("/")[-1]
    raw = unicodedata.normalize("NFKD", raw).encode("ascii", "ignore").decode("ascii")
    stem, ext = os.path.splitext(raw)
    ext = ext.lower()
    stem = _SAFE.sub("-", stem).strip("-._") or "attachment"
    if ext not in ALLOWED_EXT:
        ext = ""
    return (stem[:80] + ext) or "attachment"


def unique_in(folder: str, name: str) -> str:
    """`notes.txt` -> `notes_2.txt` when `notes.txt` is already there. NUMBERED FROM WHAT IS ON
    DISK, never from a counter - the same rule the tailored documents follow."""
    stem, ext = os.path.splitext(name)
    cand, n = name, 1
    while os.path.exists(os.path.join(folder, cand)):
        n += 1
        cand = "%s_%d%s" % (stem, n, ext)
        if n > 500:                                  # a bound, so a broken caller cannot spin here
            cand = "%s_%d%s" % (stem, int(time.time()), ext)
            break
    return cand


def _strip_timecodes(text: str) -> str:
    """A .vtt/.srt transcript is 60% timestamps. Keep the WORDS; the timing is not the content.

    Lines that are a cue number, a `00:00:01.000 --> 00:00:04.000` range, or the WEBVTT header are
    dropped; everything else is kept exactly as written, including speaker labels."""
    out, prev = [], ""
    for line in (text or "").splitlines():
        if _TIMECODE.match(line):
            continue
        s = line.strip()
        if not s or s == prev:          # a caption repeated across two cues is one sentence
            prev = s
            continue
        prev = s
        out.append(s)
    return "\n".join(out).strip()


def extract(name: str, blob: bytes) -> tuple:
    """(text, note) - the text we could read, and why there is none when there is none.

    NEVER RAISES on a broken file. A corrupt .pptx must not lose him the upload; it loses him the
    TEXT, and the note says which. That is the difference between a file we kept and a 500."""
    ext = ext_of(name)
    try:
        if ext in TEXT_EXT:
            t = blob.decode("utf-8", "replace")
            if ext in (".vtt", ".srt"):
                t = _strip_timecodes(t)
            return t.strip()[:MAX_TEXT], ""
        if ext == ".pdf":
            from pypdf import PdfReader
            rd = PdfReader(io.BytesIO(blob))
            if getattr(rd, "is_encrypted", False):
                try:
                    rd.decrypt("")
                except Exception:
                    return "", "the PDF is password-protected, so no text could be read"
            t = "\n".join((p.extract_text() or "") for p in rd.pages).strip()
            if not t:
                return "", "this PDF has no text layer (it is a scan or an image)"
            return t[:MAX_TEXT], ""
        if ext == ".docx":
            import docx
            d = docx.Document(io.BytesIO(blob))
            parts = [p.text for p in d.paragraphs]
            for tb in d.tables:
                for row in tb.rows:
                    parts.append("\t".join(c.text for c in row.cells))
            t = "\n".join(x for x in parts if x.strip()).strip()
            return (t[:MAX_TEXT], "") if t else ("", "the document is empty")
        if ext == ".pptx":
            from pptx import Presentation
            pres = Presentation(io.BytesIO(blob))
            parts = []
            for i, slide in enumerate(pres.slides, 1):
                said = []
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False):
                        for para in shape.text_frame.paragraphs:
                            line = "".join(r.text for r in para.runs).strip()
                            if line:
                                said.append(line)
                    # A TABLE on a slide carries the numbers, and they are the point of the slide.
                    if getattr(shape, "has_table", False):
                        for row in shape.table.rows:
                            said.append("\t".join(c.text for c in row.cells))
                # SPEAKER NOTES ARE THE MOST USEFUL PART of a deck somebody sends after a call.
                try:
                    if slide.has_notes_slide:
                        nt = (slide.notes_slide.notes_text_frame.text or "").strip()
                        if nt:
                            said.append("[notes] " + nt)
                except Exception:
                    pass
                if said:
                    parts.append("[slide %d]\n%s" % (i, "\n".join(said)))
            t = "\n\n".join(parts).strip()
            return (t[:MAX_TEXT], "") if t else ("", "the deck has no text on any slide")
        if ext == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
            parts = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        parts.append("\t".join(cells))
            t = "\n".join(parts).strip()
            return (t[:MAX_TEXT], "") if t else ("", "the spreadsheet has no values")
        if ext in IMAGE_EXT:
            # We do not OCR. Saying so is better than an empty box he cannot interpret.
            return "", "an image is kept as-is - there is no text extraction for pictures"
        if ext == ".doc":
            return "", "legacy .doc cannot be read - save it as .docx or PDF"
    except Exception as e:
        _log(evt="attach_extract_failed", name=safe_name(name), err=repr(e)[:160])
        return "", "the file could not be read: %s" % type(e).__name__
    return "", "no text extractor for this kind of file"


def folder_for(job_folder: str) -> str:
    p = os.path.join(job_folder, "attachments")
    os.makedirs(p, exist_ok=True)
    return p


def save(job_folder: str, name: str, blob: bytes) -> dict:
    """Write the bytes, read the text ONCE, and return what the index should record.

    The extracted text is written beside the file as `<name>.extracted.txt`, so it can be read
    later without the parser - and so he can open it himself. A write failure raises; a text
    failure never does."""
    if len(blob) > MAX_BYTES:
        raise ValueError("the file is larger than %d MB" % (MAX_BYTES // (1024 * 1024)))
    ext = ext_of(name)
    if ext not in ALLOWED_EXT:
        raise ValueError("%s files are not accepted here - allowed: %s"
                         % (ext or "extension-less", ", ".join(sorted(ALLOWED_EXT))))
    folder = folder_for(job_folder)
    fname = unique_in(folder, safe_name(name))
    path = os.path.join(folder, fname)
    # tmp + replace: a half-written attachment must never appear in the list as a real one.
    tmp = path + ".part"
    with open(tmp, "wb") as f:
        f.write(blob)
    os.replace(tmp, path)
    text, note = extract(fname, blob)
    if text:
        with open(path + ".extracted.txt", "w", encoding="utf-8") as f:
            f.write(text)
    _log(evt="attach_saved", name=fname, bytes=len(blob), chars=len(text), note=note)
    return {"name": fname, "bytes": len(blob), "kind": ext.lstrip("."),
            "chars": len(text), "words": len(text.split()), "text_note": note,
            "has_text": bool(text)}


def read_text(job_folder: str, name: str) -> str:
    """The extracted text for one attachment, or "" - never an exception, never a guess."""
    p = os.path.join(folder_for(job_folder), safe_name(name) + ".extracted.txt")
    try:
        with open(p, encoding="utf-8") as f:
            return f.read()
    except Exception:
        return ""


def remove(job_folder: str, name: str) -> bool:
    """Delete the file AND its extracted text. Scoped through safe_name, so a crafted name cannot
    reach outside the attachments folder."""
    folder = folder_for(job_folder)
    fname = safe_name(name)
    p = os.path.join(folder, fname)
    # BELT AND BRACES: prove the resolved path is still inside the folder before unlinking.
    if os.path.commonpath([os.path.abspath(folder), os.path.abspath(p)]) != os.path.abspath(folder):
        return False
    ok = False
    for cand in (p, p + ".extracted.txt"):
        try:
            os.remove(cand)
            ok = ok or cand == p
        except FileNotFoundError:
            pass
        except Exception as e:
            _log(evt="attach_error", where="remove", err=repr(e)[:160])
    return ok


def path_of(job_folder: str, name: str) -> str:
    """The absolute path to serve, or "" if the name escapes the folder or is not there."""
    folder = folder_for(job_folder)
    p = os.path.join(folder, safe_name(name))
    if os.path.commonpath([os.path.abspath(folder), os.path.abspath(p)]) != os.path.abspath(folder):
        return ""
    return p if os.path.isfile(p) else ""


def _selftest() -> int:
    import tempfile
    fails = []

    def ck(c, m):
        print(("  ok   " if c else "  FAIL ") + m)
        if not c:
            fails.append(m)

    print("[attach] contracts")
    # ---- the name is ours, not the browser's ------------------------------------------------
    ck(safe_name("../../etc/passwd") == "passwd", "a traversal is reduced to one path segment")
    ck("/" not in safe_name("a/b/c.txt") and safe_name("a/b/c.txt") == "c.txt",
       "a path separator never survives")
    ck(safe_name(r"C:\Users\feran\secret.txt") == "secret.txt",
       "a WINDOWS path is reduced too - basename() alone would keep it whole on Linux")
    ck("\x00" not in safe_name("ev\x00il.txt"), "a NUL byte is stripped before it truncates a path")
    ck(safe_name("run.sh") == "run", "a disallowed extension is dropped, not kept")
    ck(safe_name("") == "attachment" and safe_name("...") == "attachment",
       "an empty or punctuation-only name still produces something storable")
    ck(safe_name("интервью transcript.txt").endswith(".txt"),
       "a non-ascii name keeps its extension and stays a single segment")

    tmp = tempfile.mkdtemp(prefix="jhwatt")
    # ---- a collision is numbered, never overwritten -------------------------------------------
    r1 = save(tmp, "notes.txt", b"first interview, they asked about kubernetes")
    r2 = save(tmp, "notes.txt", b"second interview, different content entirely")
    ck(r1["name"] == "notes.txt" and r2["name"] == "notes_2.txt",
       "a second file with the same name is NUMBERED - an upload never destroys the first")
    ck(open(os.path.join(folder_for(tmp), "notes.txt"), "rb").read().endswith(b"kubernetes"),
       "...and the first file still holds its own bytes")

    # ---- the text is extracted once, at upload ------------------------------------------------
    ck(r1["has_text"] and "kubernetes" in read_text(tmp, "notes.txt"),
       "a text file is searchable immediately after it is uploaded")

    vtt = (b"WEBVTT\n\n1\n00:00:01.000 --> 00:00:04.000\nInterviewer: tell me about the migration."
           b"\n\n2\n00:00:04.000 --> 00:00:09.000\nEvgeny: we moved 400 VMs off VMware.\n")
    rv = save(tmp, "call.vtt", vtt)
    tv = read_text(tmp, rv["name"])
    ck("400 VMs off VMware" in tv and "-->" not in tv and "WEBVTT" not in tv,
       "a transcript keeps the WORDS and drops the timecodes")

    # A REAL PPTX, built here - a fixture that is not the real format tests the fixture.
    try:
        from pptx import Presentation
        from pptx.util import Inches
        pres = Presentation()
        s = pres.slides.add_slide(pres.slide_layouts[5])
        s.shapes.title.text = "Infrastructure Transformation"
        box = s.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2))
        box.text_frame.text = "VMware exit across 3 datacentres"
        s.notes_slide.notes_text_frame.text = "they want someone who has actually done the cutover"
        buf = io.BytesIO()
        pres.save(buf)
        rp = save(tmp, "their deck.pptx", buf.getvalue())
        tp = read_text(tmp, rp["name"])
        ck(rp["name"] == "their-deck.pptx", "a space in the name becomes a hyphen, not a new segment")
        ck("Infrastructure Transformation" in tp and "VMware exit" in tp,
           "a .pptx is read into text, slide by slide")
        ck("[notes] they want someone who has actually done the cutover" in tp,
           "the SPEAKER NOTES come across - the most useful part of a deck somebody sends")
    except ImportError:
        ck(False, "python-pptx is installed (a skip is not a pass)")

    # ---- a file with no text says WHY, and is still kept ---------------------------------------
    png = (b"\x89PNG\r\n\x1a\n" + b"\x00" * 40)
    ri = save(tmp, "whiteboard.png", png)
    ck(not ri["has_text"] and "no text extraction for pictures" in ri["text_note"],
       "an image is kept and the reason there is no text is RECORDED, not left blank")
    ck(os.path.isfile(os.path.join(folder_for(tmp), "whiteboard.png")),
       "...and the file itself is still on disk - it is his")
    # pypdf logs its own complaint about the corrupt fixture. A scary line printed on EVERY run
    # trains you to read past the log, so the deliberate one is silenced here and only here.
    import logging as _lg
    _lg.getLogger("pypdf").setLevel(_lg.CRITICAL)
    rb = save(tmp, "broken.pdf", b"this is not a pdf at all")
    ck(not rb["has_text"] and rb["text_note"],
       "a corrupt file loses its TEXT, never the upload - and the note says so")

    # ---- what is refused ----------------------------------------------------------------------
    try:
        save(tmp, "payload.exe", b"MZ")
        ck(False, "an executable is refused")
    except ValueError as e:
        ck("not accepted" in str(e), "an executable is refused, by allowlist, with the reason")
    try:
        save(tmp, "huge.txt", b"x" * (MAX_BYTES + 1))
        ck(False, "an oversized file is refused")
    except ValueError as e:
        ck("larger than" in str(e), "an oversized file is refused before it is written")
    ck(not os.path.exists(os.path.join(folder_for(tmp), "huge.txt")),
       "...and a refused upload leaves NOTHING on disk")
    ck(not [f for f in os.listdir(folder_for(tmp)) if f.endswith(".part")],
       "no half-written .part file survives a save")

    # ---- reach ---------------------------------------------------------------------------------
    ck(path_of(tmp, "notes.txt").endswith("notes.txt"), "a real attachment resolves to a path")
    ck(path_of(tmp, "../../../etc/passwd") == "",
       "a crafted name cannot resolve to a file outside the folder")
    ck(remove(tmp, "../../../etc/passwd") is False, "...and cannot delete one either")
    ck(remove(tmp, "notes.txt") is True and path_of(tmp, "notes.txt") == "",
       "his own attachment deletes")
    ck(not os.path.exists(os.path.join(folder_for(tmp), "notes.txt.extracted.txt")),
       "...and its extracted text goes with it, so nothing readable is left behind")

    print("attach selftest: %d check(s) failed" % len(fails))
    return 1 if fails else 0


if __name__ == "__main__":
    import sys
    sys.exit(_selftest())
